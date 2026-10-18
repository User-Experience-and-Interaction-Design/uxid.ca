"""Build venues-data.json for venues.html from the deadline workbook.

The spreadsheet is the single source of truth; this only reshapes it for the
browser. Run it after editing the workbook:

    python3 tools/export-venues.py

Needs openpyxl (pip install openpyxl). Override the paths with the
VENUES_XLSX and VENUES_JSON environment variables if the workbook moves.
"""
import openpyxl, json, datetime, re, os

DEFAULT_SRC = os.path.expanduser(
    "~/Library/Mobile Documents/com~apple~CloudDocs/Conference Deadlines/"
    "Conf new/HCI_Conferences_Journals_2026-27_v10.xlsx")
SRC = os.environ.get("VENUES_XLSX", DEFAULT_SRC)
OUT = os.environ.get(
    "VENUES_JSON",
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                 "venues-data.json"))

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["HCI Conferences & Journals"]
hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
ix = {h: i for i, h in enumerate(hdr) if h}

# The workbook already sorts every track into a category sheet, so take the
# category from there rather than guessing it from the Submission_Type text.
CAT_SHEETS = ["Full Papers", "Short Papers", "Posters", "Late-Breaking Work",
              "Work in Progress", "Demos", "Workshops", "Doctoral",
              "Tutorials & Courses", "Industry", "Competitions",
              "Panels & SIGs", "Journals", "Other"]
catof = {}
for sheet in CAT_SHEETS:
    if sheet not in wb.sheetnames:
        continue
    t = wb[sheet]
    for r in range(2, t.max_row + 1):
        nm = t.cell(r, ix["Name"] + 1).value
        if not nm:
            continue
        key = (str(nm), str(t.cell(r, ix["Submission_Type"] + 1).value),
               str(t.cell(r, ix["Deadline"] + 1).value))
        catof.setdefault(key, sheet)

def iso(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s): return s
    return None

def txt(v):
    if v is None: return ""
    s = str(v).strip()
    return "" if s.upper() in ("N/A", "NA", "NONE", "TBD", "-") else s

TAGS = [("Tag_Mobile","Mobile"),("Tag_Smartwatch","Smartwatch"),("Tag_HumanAI","Human-AI"),
        ("Tag_SystemDesign","Systems & design"),("Tag_Visualization","Visualization"),
        ("Tag_Accessibility","Accessibility"),("Tag_Health","Health"),("Tag_Education","Education"),
        ("Tag_CSCW","CSCW"),("Tag_Privacy_Security","Privacy"),("Tag_Robotics","Robotics"),
        ("Tag_NLP_Voice","NLP & voice"),("Tag_Ethics_Fairness","Ethics")]

# ---- HCI relevance, normalised from 10 free-text variants in the sheet ----
def relbucket(v):
    t = v.lower()
    if t.startswith("core"):                      return "Core"
    if "high" in t or "specialized" in t:         return "High"
    if t.startswith("adjacent"):                  return "Adjacent"
    if t.startswith("indirect"):                  return "Indirect"
    return "Adjacent"

# ---- attendance, from Presentation_Mode ----
def attend(pm, fmt):
    t = (pm or "").lower()
    if "virtual event" in t:          return "Virtual"
    if "remote option" in t:          return "Remote option"
    if "by exception" in t:           return "Remote by exception"
    if "in-person required" in t:     return "In person only"
    if "no presentation" in t:        return "No presentation"
    return fmt or ""

# ---- Fit to the UXID Lab's own areas -------------------------------------
# The workbook has a Fit_Score column but it is empty, so this is computed.
# Weights reflect the lab's stated research: small-screen and smartwatch
# interaction and visualization, wearables, human-AI, accessibility.
# Change the weights here and re-export; nothing else needs touching.
AREA_W = {"Smartwatch": 1.0, "Mobile": 0.8, "Human-AI": 0.8,
          "Visualization": 0.6, "Accessibility": 0.5}
REL_W  = {"Core": 1.2, "High": 1.0, "Adjacent": 0.5, "Indirect": 0.0}
# Tier MULTIPLIES rather than adds. Added, it let a weak venue tagged across
# many areas (ACHI, Tier 3, tagged Yes on seven) outrank UbiComp/IMWUT.
TIER_M = {1: 1.0, 2: 0.85, 3: 0.55, 4: 0.35}

def fitraw(yes, part, rel, tier):
    topic = REL_W.get(rel, 0.4)
    for a, w in AREA_W.items():
        if a in yes:    topic += w
        elif a in part: topic += w / 2.0
    return topic * TIER_M.get(tier, 0.5)

rows, skipped = [], 0
for r in range(2, ws.max_row + 1):
    g = lambda k: ws.cell(r, ix[k] + 1).value
    if not g("Name"):
        continue
    raw = str(g("Deadline") or "").strip()
    dl = iso(g("Deadline"))
    # Rolling journals and not-yet-announced tracks are real options and stay;
    # discontinued venues are dropped.
    if "discontinu" in raw.lower():
        skipped += 1
        continue
    st = "dated" if dl else ("rolling" if "rolling" in raw.lower() else "tba")
    tier = txt(g("Tier"))
    m = re.search(r"\d", tier)
    tier_n = int(m.group()) if m else 0
    yes  = [lbl for col, lbl in TAGS if txt(g(col)) == "Yes"]
    part = [lbl for col, lbl in TAGS if txt(g(col)) == "Partial"]
    rel  = relbucket(txt(g("HCI_Relevance")))
    try:
        sheet_fit = float(g("Fit_Score")) if g("Fit_Score") not in (None, "") else None
    except (TypeError, ValueError):
        sheet_fit = None
    cat = catof.get((str(g("Name")), str(g("Submission_Type")), str(g("Deadline"))), "Other")
    rows.append({
        "cat":  cat,
        "ty":   "J" if txt(g("Type")).lower().startswith("j") else "C",
        "n":    txt(g("Name")),
        "fn":   txt(g("Full_Name")),
        "tier": tier_n,
        "tr":   txt(g("Submission_Type")),
        "dl":   dl,
        "st":   st,
        "dlraw": "" if dl else raw,
        "ab":   iso(g("Abstract_Deadline")),
        "nt":   iso(g("Notification_Date")),
        "cr":   iso(g("Camera_Ready")),
        "us":   txt(g("User_Study_Policy")),
        "pg":   txt(g("Page_Limit")),
        "acc":  txt(g("Acceptance_Rate")),
        "core": txt(g("CORE_Ranking")),
        "ccf":  txt(g("CCF_Ranking")),
        "loc":  txt(g("Location")),
        "cd":   txt(g("Conference_Dates")),
        "web":  txt(g("Website")),
        "rel":  txt(g("HCI_Relevance")),
        "why":  txt(g("Why_HCI_Relevant")),
        # keep only the areas the venue actually serves
        "yes":  yes,
        "part": part,
        "relb": rel,
        "fitraw": fitraw(yes, part, rel, tier_n),
        "fitsheet": sheet_fit,
        "att":  attend(txt(g("Presentation_Mode")), txt(g("Format"))),
        "fmt":  txt(g("Format")),
        "pm":   txt(g("Presentation_Mode")),
        "rod":  txt(g("Remote_Option_Details")),
        "size": txt(g("Venue_Size_Annual")),
        "subs": txt(g("Annual_Submissions")),
        "reg":  txt(g("Registration_Requirement")),
    })
    # Journal-only fields: metrics and the practical "how do I submit" detail.
    if rows[-1]["ty"] == "J":
        hm = re.search(r"\d+", str(g("Venue_H_Index") or ""))
        rows[-1].update({
            "pub":   txt(g("Publisher")),
            "if_":   txt(g("Impact_Factor")),
            "cs":    txt(g("CiteScore_2023")),
            "h":     int(hm.group()) if hm else None,
            "sys":   txt(g("Submission_System")),
            "portal": txt(g("Submission_Portal")),
            "tmpl":  txt(g("Template")),
            "rev":   txt(g("Review_Process")),
            "turn":  txt(g("Review_Turnaround")),
            "nrev":  txt(g("Reviewer_Count")),
            "oa":    txt(g("Open_Access_Policy")),
            "apc":   txt(g("APC_Submission_Fee")),
            "idx":   txt(g("Indexing_Databases")),
            "freq":  txt(g("Frequency")),
            "sip":   txt(g("Special_Issue_Partner")),
        })

# Normalise to 0.5-5 against the closest-fitting venue in the list, so the
# top of the scale is actually used. 5 means "the best fit here", not perfect.
peak = max(x["fitraw"] for x in rows) or 1.0
for x in rows:
    v = x.pop("fitsheet")
    x["fit"] = v if v else max(0.5, round(x.pop("fitraw") / peak * 5 * 2) / 2.0)
    x.pop("fitraw", None)

rows.sort(key=lambda x: (x["dl"] or "9999", x["n"], x["tr"]))
areas = [lbl for _, lbl in TAGS]
chi_h = None
for r in range(2, ws.max_row + 1):
    if ws.cell(r, ix["Name"] + 1).value == "CHI":
        m = re.search(r"\d+", str(ws.cell(r, ix["Venue_H_Index"] + 1).value or ""))
        if m:
            chi_h = int(m.group())
        break
payload = {"generated": datetime.date.today().isoformat(), "source": "HCI_Conferences_Journals_2026-27_v10.xlsx",
           "areas": areas, "chi_h": chi_h, "rows": rows}
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))

import os
import collections
print("by status:", dict(collections.Counter(x["st"] for x in rows)))
print("by category:")
for k, v in collections.Counter(x["cat"] for x in rows).most_common():
    print(f"   {v:4}  {k}")
print("fit distribution:", dict(sorted(collections.Counter(x["fit"] for x in rows).items())))
best = {}
for x in rows:
    best[x["n"]] = max(best.get(x["n"], 0), x["fit"])
print("top venues by fit:", sorted(best.items(), key=lambda kv: -kv[1])[:14])
print("journals:", sum(1 for x in rows if x["ty"] == "J"), "| CHI h5 =", chi_h)
print("rows:", len(rows), "| dropped (discontinued):", skipped,
      "| venues:", len({x['n'] for x in rows}),
      "| bytes:", os.path.getsize(OUT))
